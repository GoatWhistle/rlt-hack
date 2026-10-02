"""Проверки полного XLSX-снимка ГИСП без сети и реального реестра."""

import asyncio
import io
import sys
from datetime import datetime
from pathlib import Path

import httpx
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from src.adapter.supplier import identity
from src.adapter.supplier.errors import ContentFormatError, SourceUnavailableError
from src.adapter.supplier.gisp_registry import GispRegistryProvider
from src.models.catalog.source import Source
from src.models.enums import SourceType, SupplierRole

URL = "https://gisp.gov.ru/pp719v2/pub/prod/"
EXPORT = "https://gisp.gov.ru/export.xlsx"
SOURCE = Source(
    identity.source_id(URL, "gisp_registry"), "ГИСП", URL, SourceType.REGISTRY, "gisp_registry"
)
HEADERS = [
    "Реестровый номер",
    "Наименование продукции",
    "Наименование организации",
    "ИНН",
    "ОКПД2",
    "Статус",
    "Срок действия",
]
ROWS = [
    [
        "10001",
        "Станок токарный",
        "ООО Завод",
        "7707083893",
        "28.41.21.110",
        "Действует",
        "2027-12-31",
    ],
    ["10002", "Станок фрезерный", "ООО Завод", "7707083893", "28.41.22.110", "Истёк", "2025-12-31"],
]


def workbook(rows: list[list[str]]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def grouped_workbook() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["Предприятие", None, "Продукция", None, None])
    sheet.append(["Наименование", "ОГРН", "Реестровый номер", "Наименование", "ОКПД2"])
    sheet.append(["ООО Завод", "1027700132195", "10001", "Станок токарный", "28.41.21.110"])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def official_layout_workbook() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["Время выгрузки: 01.10.2026, 01:00:10"])
    sheet.append([])
    sheet.append(
        [
            "Предприятие",
            "ИНН",
            "ОГРН",
            "Фактический адрес производителя",
            "Адрес производственных помещений",
            "Первичный регистрационный номер реестровой записи",
            "Реестровый номер",
            "Номер цифрового паспорта",
            "Дата внесения в реестр",
            "Срок действия",
            "Фактическая дата прекращения действия реестровой записи",
            "Наименование продукции",
            "ОКПД2",
            "ТН ВЭД",
            "Изготовлена по",
            "Баллы",
            "Процентный показатель",
            "О соответствии",
            "Искусственный интеллект",
            "Высокотехнологичное оборудование",
            "Доверенный ПАК",
            "Основание: Наименование",
            "Основание: Дата",
            "Основание: Номер",
            "Основание: Срок действия",
            "Заключение: Департамент",
            "Заключение: Номер заключения",
            "Заключение: Документ",
        ]
    )
    row = [None] * 28
    row[0] = "ООО Завод"
    row[1] = "7707083893"
    row[2] = "1027700132195"
    row[3] = "Москва, улица Заводская"
    row[6] = "10902841"
    row[8] = datetime(2026, 10, 1)
    row[9] = datetime(2029, 10, 1)
    row[11] = "Станок токарный"
    row[12] = "28.41.21.110"
    row[13] = "8466 10 380 0"
    row[21] = "СТ-1"
    row[23] = "6036000122"
    sheet.append(row)
    expired = row.copy()
    expired[6] = "10902842"
    expired[10] = datetime(2026, 9, 1)
    expired[11] = "Станок фрезерный"
    sheet.append(expired)
    version = row.copy()
    version[8] = datetime(2025, 10, 1)
    version[10] = datetime(2025, 11, 1)
    version[23] = "6036000123"
    sheet.append(version)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def provider(content: bytes, status: int = 200) -> GispRegistryProvider:
    def response(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/org/b/"):
            return httpx.Response(
                200,
                json={
                    "ok": True,
                    "items": [
                        {
                            "org_name": "ООО Без продукции",
                            "org_inn": "4707019370",
                            "gisp_url": "https://gisp.gov.ru/company-catalog/company/2/",
                        }
                    ],
                    "total_count": 1,
                },
            )
        return httpx.Response(status, content=content)

    transport = httpx.MockTransport(response)
    return GispRegistryProvider(SOURCE, EXPORT, transport=transport)


async def merge_checks() -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(
        [
            "Реестровый номер",
            "Наименование продукции",
            "Наименование организации",
            "ИНН",
            "КПП",
            "ОГРН",
            "Сайт",
            "Фактический адрес производителя",
        ]
    )
    sheet.append(
        [
            "10001",
            "Станок",
            "ООО Из выгрузки",
            "7707083893",
            "770701001",
            "1027700132195",
            "https://factory.example",
            "Москва, улица Заводская",
        ]
    )
    output = io.BytesIO()
    book.save(output)

    def response(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/org/b/"):
            return httpx.Response(
                200,
                json={
                    "ok": True,
                    "items": [
                        {
                            "org_name": "ООО Из перечня",
                            "org_inn": "7707083893",
                            "org_ogrn": "1027700000000",
                            "org_region_name": "Москва",
                            "gisp_url": "https://gisp.gov.ru/company-catalog/company/1/",
                        }
                    ],
                    "total_count": 1,
                },
            )
        return httpx.Response(200, content=output.getvalue())

    package = await GispRegistryProvider(
        SOURCE, EXPORT, transport=httpx.MockTransport(response)
    ).fetch()
    supplier = package.suppliers[0]
    assert supplier.name == "ООО Из перечня"
    assert supplier.kpps == ("770701001",)
    assert supplier.website == "https://factory.example"
    assert supplier.region == "Москва"
    assert supplier.contacts["ogrn"] == "1027700000000"
    assert supplier.contacts["actual_address"] == "Москва, улица Заводская"
    assert supplier.identity_evidence_url == "https://gisp.gov.ru/company-catalog/company/1/"
    assert package.offers[0].supplier_id == supplier.supplier_id


async def checks() -> None:
    await merge_checks()
    first = await provider(workbook(ROWS)).fetch()
    again = await provider(workbook(ROWS)).fetch()
    assert len(first.suppliers) == 2
    assert len(first.offers) == 2
    assert {offer.supplier_id for offer in first.offers} == {first.suppliers[1].supplier_id}
    assert {offer.offer_id for offer in first.offers} == {offer.offer_id for offer in again.offers}
    assert {offer.content_hash for offer in first.offers} == {
        offer.content_hash for offer in again.offers
    }
    assert first.offers[0].supplier_role == SupplierRole.MANUFACTURER
    assert first.offers[1].supplier_role == SupplierRole.UNKNOWN
    assert first.offers[0].price is None
    assert first.offers[0].attributes["registry_status"] == "Действует"
    assert first.offers[1].attributes["registry_status"] == "Истёк"
    changed = [ROWS[0].copy(), ROWS[1].copy()]
    changed[0][5] = "Истёк"
    after = await provider(workbook(changed)).fetch()
    assert after.offers[0].offer_id == first.offers[0].offer_id
    assert after.offers[0].supplier_role == SupplierRole.UNKNOWN
    assert after.offers[0].content_hash != first.offers[0].content_hash
    grouped = await provider(grouped_workbook()).fetch()
    assert len(grouped.suppliers) == 2
    assert grouped.offers[0].name == "Станок токарный"
    assert grouped.suppliers[1].inn is None
    unidentified = ROWS[0].copy()
    unidentified[3] = ""
    unknown_company = await provider(workbook([unidentified])).fetch()
    assert unknown_company.suppliers[1].inn is None
    assert unknown_company.offers[0].supplier_id == unknown_company.suppliers[1].supplier_id
    official = await provider(official_layout_workbook()).fetch()
    assert len(official.offers) == 3
    assert official.offers[0].supplier_role == SupplierRole.MANUFACTURER
    assert official.offers[1].supplier_role == SupplierRole.UNKNOWN
    assert official.offers[2].supplier_role == SupplierRole.UNKNOWN
    assert (
        official.offers[0].attributes["registry_number"]
        == official.offers[2].attributes["registry_number"]
    )
    assert official.offers[0].offer_id != official.offers[2].offer_id
    assert official.offers[0].attributes["basis_number"] == "6036000122"
    assert official.offers[0].attributes["registry_introduced"] == "2026-10-01"
    assert official.suppliers[1].contacts["actual_address"] == "Москва, улица Заводская"
    invalid_row = ["10003", "", "ООО Завод", "7707083893"]
    for content in (workbook([]), b"not xlsx", workbook([invalid_row])):
        try:
            await provider(content).fetch()
        except ContentFormatError:
            pass
        else:
            raise AssertionError("Неполный или неверный экспорт не должен сохраняться")
    try:
        await provider(workbook(ROWS), 503).fetch()
    except SourceUnavailableError:
        pass
    else:
        raise AssertionError("Сбой HTTP не должен сохранять снимок")
    duplicates = await provider(workbook([ROWS[0], ROWS[0]])).fetch()
    assert len(duplicates.offers) == 1


if __name__ == "__main__":
    asyncio.run(checks())
    print("ГИСП: проверки прошли")
