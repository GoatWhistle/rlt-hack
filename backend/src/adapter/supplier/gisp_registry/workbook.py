"""Разбор строк полного экспорта реестра; неизвестная схема прерывает обход."""

import io
import re
from datetime import UTC, datetime

from openpyxl import load_workbook

from src.adapter.supplier import identity
from src.adapter.supplier.errors import ContentFormatError
from src.adapter.supplier.gisp_registry.record import active_record, registry_external_id
from src.adapter.supplier.inn import normalize_inn, normalize_kpp
from src.models.enums import ItemType, SupplierRole, VerificationStatus
from src.models.offer import Offer
from src.models.package import SupplierPackage
from src.models.source import Source
from src.models.supplier import Supplier

_FIELDS = {
    "number": (
        "реестровый номер",
        "номер реестровой записи",
        "номер записи",
        "продукция реестровый номер",
    ),
    "product": (
        "наименование продукции",
        "наименование производимой промышленной продукции",
        "продукция наименование",
    ),
    "company": (
        "предприятие",
        "наименование предприятия",
        "наименование организации",
        "производитель",
        "предприятие наименование",
    ),
    "inn": ("инн", "инн производителя", "инн организации"),
    "ogrn": ("огрн", "огрн производителя", "огрн организации", "предприятие огрн"),
    "kpp": ("кпп", "кпп организации"),
    "okpd2": ("окпд2", "код окпд2", "код промышленной продукции по окпд2"),
    "tnved": ("тн вэд", "код тн вэд", "тн вэд еаэс"),
    "status": ("статус", "статус записи", "состояние"),
    "expires": ("срок действия", "дата окончания действия", "действует до"),
    "ceased": ("фактическая дата прекращения действия реестровой записи",),
    "introduced": ("дата внесения в реестр",),
    "initial_number": ("первичный регистрационный номер реестровой записи",),
    "passport_number": ("номер цифрового паспорта",),
    "document": ("ссылка на документ", "документ", "ссылка на выписку"),
    "address": ("фактический адрес производителя",),
    "facilities": ("адрес производственных помещений",),
    "region": ("регион", "субъект рф"),
    "website": ("сайт", "веб сайт"),
    "description": ("описание продукции", "описание"),
    "standard": ("изготовлена по", "нормативный документ"),
    "points": ("баллы", "количество баллов"),
    "percent": ("процентный показатель",),
    "compliance": ("о соответствии",),
    "ai": ("искусственный интеллект",),
    "hightech": ("высокотехнологичное оборудование",),
    "trusted_pak": ("доверенный пак",),
    "basis_name": ("основание наименование",),
    "basis_date": ("основание дата",),
    "basis_number": ("основание номер",),
    "basis_expires": ("основание срок действия",),
    "department": ("заключение департамент",),
    "conclusion_number": ("заключение номер заключения",),
    "conclusion_document": ("заключение документ",),
}
_REQUIRED = ("number", "product", "company")


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    return " ".join(str(value).split())


def _key(value: object) -> str:
    return re.sub(r"[^\w]+", " ", _text(value).casefold()).strip()


def _columns(row: tuple[object, ...], parent: tuple[object, ...] = ()) -> dict[str, int]:
    found: dict[str, int] = {}
    group = ""
    for index, value in enumerate(row):
        label = _key(value)
        if index < len(parent) and _key(parent[index]) in ("предприятие", "продукция"):
            group = _key(parent[index])
        combined = f"{group} {label}" if group and label else ""
        for field, aliases in _FIELDS.items():
            if label in aliases or combined in aliases:
                found[field] = index
    return found


def _value(row: tuple[object, ...], columns: dict[str, int], field: str) -> str:
    index = columns.get(field)
    return _text(row[index]) if index is not None and index < len(row) else ""


def parse_workbook(content: bytes, source: Source) -> SupplierPackage:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ContentFormatError(f"Ответ не является XLSX: {error}") from error
    suppliers: dict[str, Supplier] = {}
    offers: dict[str, Offer] = {}
    found_header = False
    observed_at = datetime.now(UTC)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            columns: dict[str, int] = {}
            previous: tuple[object, ...] = ()
            for row_number, row in enumerate(rows, start=1):
                if not columns:
                    columns = _columns(row, previous)
                    if all(field in columns for field in _REQUIRED):
                        found_header = True
                    else:
                        columns = {}
                    previous = row
                    if row_number >= 30 and not found_header:
                        break
                    continue
                if not any(value is not None for value in row):
                    continue
                number = _value(row, columns, "number")
                product = _value(row, columns, "product")
                company = _value(row, columns, "company")
                if not (number and product and company):
                    raise ContentFormatError(f"{sheet.title}:{row_number}: неполная запись")
                inn = normalize_inn(_value(row, columns, "inn"))
                ogrn = _value(row, columns, "ogrn")
                company_key = inn or ogrn or f"registry:{number}"
                supplier_uuid = identity.supplier_id(inn, source.source_id, company_key)
                kpp = normalize_kpp(_value(row, columns, "kpp"))
                supplier = Supplier(
                    supplier_id=supplier_uuid,
                    name=company,
                    inn=inn,
                    kpps=(kpp,) if kpp else (),
                    region=_value(row, columns, "region"),
                    website=_value(row, columns, "website"),
                    contacts={
                        key: value
                        for key, value in (
                            ("ogrn", ogrn),
                            ("actual_address", _value(row, columns, "address")),
                            ("production_address", _value(row, columns, "facilities")),
                        )
                        if value
                    },
                    identity_status=VerificationStatus.UNVERIFIED,
                    identity_evidence_url=source.base_url,
                )
                suppliers[company_key] = supplier
                status = _value(row, columns, "status")
                expires = _value(row, columns, "expires")
                ceased = _value(row, columns, "ceased")
                active = active_record(status, expires, ceased, observed_at.date())
                attributes = {
                    key: value
                    for key, value in (
                        ("registry_number", number),
                        ("registry_status", status),
                        ("registry_expires", expires),
                        ("registry_ceased", ceased),
                        ("registry_introduced", _value(row, columns, "introduced")),
                        ("initial_registry_number", _value(row, columns, "initial_number")),
                        ("digital_passport_number", _value(row, columns, "passport_number")),
                        ("registry_document", _value(row, columns, "document")),
                        ("manufacturer_address", _value(row, columns, "address")),
                        ("production_address", _value(row, columns, "facilities")),
                        ("tnved", _value(row, columns, "tnved")),
                        ("standard", _value(row, columns, "standard")),
                        ("points", _value(row, columns, "points")),
                        ("percent", _value(row, columns, "percent")),
                        ("compliance", _value(row, columns, "compliance")),
                        ("ai", _value(row, columns, "ai")),
                        ("hightech", _value(row, columns, "hightech")),
                        ("trusted_pak", _value(row, columns, "trusted_pak")),
                        ("basis_name", _value(row, columns, "basis_name")),
                        ("basis_date", _value(row, columns, "basis_date")),
                        ("basis_number", _value(row, columns, "basis_number")),
                        ("basis_expires", _value(row, columns, "basis_expires")),
                        ("department", _value(row, columns, "department")),
                        ("conclusion_number", _value(row, columns, "conclusion_number")),
                        ("conclusion_document", _value(row, columns, "conclusion_document")),
                    )
                    if value
                }
                external_id = registry_external_id(
                    number,
                    _value(row, columns, "introduced"),
                    _value(row, columns, "basis_number"),
                    company_key,
                    product,
                )
                url = source.base_url
                offer = Offer(
                    offer_id=identity.offer_id(source.source_id, external_id),
                    source_id=source.source_id,
                    external_id=external_id,
                    url=url,
                    name=product,
                    first_seen_at=observed_at,
                    last_seen_at=observed_at,
                    supplier_id=supplier_uuid,
                    seller_status=VerificationStatus.UNVERIFIED,
                    evidence_url=source.base_url,
                    description=_value(row, columns, "description"),
                    item_type=ItemType.GOODS,
                    okpd2_code=_value(row, columns, "okpd2"),
                    attributes=attributes,
                    supplier_role=SupplierRole.MANUFACTURER if active else SupplierRole.UNKNOWN,
                    role_evidence_text=(
                        "Действующая запись реестра российской промышленной продукции"
                        if active
                        else ""
                    ),
                    content_hash=identity.offer_content_hash(
                        name=product,
                        description=_value(row, columns, "description"),
                        item_type=str(ItemType.GOODS),
                        okpd2_code=_value(row, columns, "okpd2"),
                        attributes=attributes,
                    ),
                )
                previous = offers.get(external_id)
                if previous and previous.content_hash != offer.content_hash:
                    raise ContentFormatError(
                        f"{sheet.title}:{row_number}: противоречивый дубль {number}"
                    )
                offers[external_id] = offer
    finally:
        workbook.close()
    if not found_header or not offers:
        raise ContentFormatError("В выгрузке нет таблицы реестровых записей")
    return SupplierPackage(
        source=source,
        suppliers=tuple(suppliers.values()),
        offers=tuple(offers.values()),
    )
